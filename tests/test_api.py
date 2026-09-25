"""Test API Overlord: cicli end-to-end over rotte, con file isolati in tmp_path."""
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent

sys.path.insert(0, str(ROOT / "app"))


@pytest.fixture()
def app(tmp_path: Path, monkeypatch):
    import app as appmod

    # Reindirizza cartelle/config/state su una dir temporanea isolata.
    files = tmp_path / "files"
    (files / "source").mkdir(parents=True)
    (files / "dest").mkdir()
    (files / "export").mkdir()
    (files / "config").mkdir()
    (files / "state").mkdir()

    monkeypatch.setattr(appmod, "SOURCE_DIR", files / "source")
    monkeypatch.setattr(appmod, "DEST_DIR", files / "dest")
    monkeypatch.setattr(appmod, "EXPORT_DIR", files / "export")
    monkeypatch.setattr(appmod, "CONFIG_PATH", files / "config" / "selezione.json")
    monkeypatch.setattr(appmod, "STATE_PATH", files / "state" / "last_value.json")

    # Mockup minimi (1 sorgente, 1 scheda).
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "VELOCITA'"
    ws["A1"].font = Font(bold=True)
    ws["A2"], ws["A3"] = 100, 101
    wb.save(files / "source" / "imp.xlsx")

    wb = Workbook()
    ws = wb.active
    ws["B3"] = "VELOCITA'"
    ws["B3"].font = Font(bold=True)
    # valore vecchio sotto l'etichetta: deve essere svuotato alla compilazione
    ws["B4"], ws["B5"], ws["B6"] = 99, 98, 97
    wb.save(files / "dest" / "scheda.xlsx")

    appmod.app.config["TESTING"] = True
    return appmod


@pytest.fixture()
def client(app):
    return app.app.test_client()


def _get(client, url):
    return client.get(url)


def test_sorgenti_destinazioni(client):
    assert _get(client, "/api/sorgenti").get_json() == ["imp.xlsx"]
    assert _get(client, "/api/destinazioni").get_json() == ["scheda.xlsx"]


def test_argomenti_e_404(client):
    r = _get(client, "/api/argomenti?file=imp.xlsx")
    assert r.status_code == 200
    data = r.get_json()
    assert data[0]["etichetta"] == "VELOCITA'"
    assert data[0]["valori"] == [100, 101]
    assert _get(client, "/api/argomenti?file=x.xlsx").status_code == 404


def test_scheda_voci(client):
    data = _get(client, "/api/scheda?file=scheda.xlsx").get_json()
    assert any(v["etichetta"] == "VELOCITA'" and v["posizione"] == "B3" for v in data)


def test_aggiungi_arriva_in_config(client):
    r = client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    assert r.status_code == 200
    assert r.get_json()["argomenti"][0]["riga_etichetta"] == 1


def test_aggiungi_duplicato_ignorato(client):
    payload = {"sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1}
    client.post("/api/aggiungi", json=payload)
    client.post("/api/aggiungi", json=payload)
    cfg = _get(client, "/api/config").get_json()
    assert len(cfg["argomenti"]) == 1


def test_aggiungi_riga_invalida(client):
    r = client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": "abc",
    })
    assert r.status_code == 400


def test_compila_svuota_e_riscrive(app, client):
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    client.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})
    r = client.post("/api/compila")
    assert r.status_code == 200
    esiti = r.get_json()["esiti"]
    assert esiti[0]["trovata"] is True
    assert esiti[0]["scritti"] == 2

    # Verifica fisica sul file isolato: valori nuovi sotto B3, vecchi svuotati.
    import openpyxl

    wb = openpyxl.load_workbook(app.DEST_DIR / "scheda.xlsx")
    ws = wb.active
    assert ws["B3"].value == "VELOCITA'"
    assert ws["B4"].value == 100
    assert ws["B5"].value == 101
    assert ws["B6"].value is None
    wb.close()


def test_export_traversal_bloccato(app, client):
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    r = client.post("/api/export", json={"nome": "../../fuga.xlsx"})
    assert r.status_code == 200
    f = r.get_json()["file"]
    assert "../" not in f
    assert "/" not in f
    assert (app.EXPORT_DIR / f).exists()


def test_nome_foglio_invalido_sanificato(client):
    # etichetta con carattere vietato in un foglio Excel: non deve rompere l'export
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VEL/OCITA'", "riga_etichetta": 1,
    })
    r = client.post("/api/export", json={"nome": "ok.xlsx"})
    assert r.status_code == 200


def test_export_download(client):
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    f = client.post("/api/export", json={"nome": "out.xlsx"}).get_json()["file"]
    r = _get(client, f"/api/download/{f}")
    assert r.status_code == 200


def test_download_sconosciuto(client):
    assert _get(client, "/api/download/nope.xlsx").status_code == 404


def test_elabora_prima_volta_poi_nessuna_variazione(client):
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    client.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})
    r1 = client.post("/api/elabora")
    assert r1.status_code == 200
    assert r1.get_json()["ok"] is True
    r2 = client.post("/api/elabora")
    assert r2.get_json()["ok"] is False
    assert "nessuna variazione" in r2.get_json()["motivo"]


def test_elabora_config_incompleta(client):
    r = client.post("/api/elabora")
    assert r.status_code == 200
    assert r.get_json()["ok"] is False


def test_rimuovi(client):
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    r = client.post("/api/rimuovi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'",
    })
    assert r.get_json()["argomenti"] == []


def test_modella_crea_voce(client):
    r = client.post("/api/modella", json={
        "destinazione": "scheda.xlsx", "testo": "FOTOMETRIA", "riga": 28, "colonna": "B",
    })
    assert r.status_code == 200
    data = _get(client, "/api/scheda?file=scheda.xlsx").get_json()
    assert any(v["etichetta"] == "FOTOMETRIA" and v["posizione"] == "B28" for v in data)


def test_modella_colonna_invalida(client):
    r = client.post("/api/modella", json={
        "destinazione": "scheda.xlsx", "testo": "X", "riga": 1, "colonna": "1",
    })
    assert r.status_code == 400


@pytest.fixture()
def app2(tmp_path: Path, monkeypatch):
    """Variante con 2 sorgenti (uno con date) e scheda piu' stretta."""
    import app as appmod

    files = tmp_path / "files2"
    (files / "source").mkdir(parents=True)
    (files / "dest").mkdir()
    (files / "export").mkdir()
    (files / "config").mkdir()
    (files / "state").mkdir()

    monkeypatch.setattr(appmod, "SOURCE_DIR", files / "source")
    monkeypatch.setattr(appmod, "DEST_DIR", files / "dest")
    monkeypatch.setattr(appmod, "EXPORT_DIR", files / "export")
    monkeypatch.setattr(appmod, "CONFIG_PATH", files / "config" / "selezione.json")
    monkeypatch.setattr(appmod, "STATE_PATH", files / "state" / "last_value.json")

    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Font

    # sorgente A: ALFA con 5 valori (troppi per la scheda stretta)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ALFA"
    ws["A1"].font = Font(bold=True)
    for i, v in enumerate([1, 2, 3, 4, 5], start=2):
        ws.cell(row=i, column=1, value=v)
    wb.save(files / "source" / "a.xlsx")

    # sorgente DATA con una cella datetime
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "DATA"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = datetime(2026, 9, 1, 10, 30)
    wb.save(files / "source" / "data.xlsx")

    # scheda: ALFA con spazio per 1 solo valore, poi etichetta BETA subito sotto
    wb = Workbook()
    ws = wb.active
    ws["B2"] = "ALFA"
    ws["B2"].font = Font(bold=True)
    ws["B3"] = 1
    ws["B5"] = "BETA"
    ws["B5"].font = Font(bold=True)
    ws["E7"] = "DATA"
    ws["E7"].font = Font(bold=True)
    wb.save(files / "dest" / "scheda.xlsx")

    appmod.app.config["TESTING"] = True
    return appmod


@pytest.fixture()
def client2(app2):
    return app2.app.test_client()


def test_argomenti_path_traversal_negato(client2):
    """Finding 4: niente lettura fuori da SOURCE_DIR da /api/argomenti."""
    r = client2.get("/api/argomenti?file=../../segreto.xlsx")
    assert r.status_code == 404


def test_scheda_path_traversal_negato(client2):
    """Finding 4: niente lettura fuori da DEST_DIR da /api/scheda."""
    r = client2.get("/api/scheda?file=../../segreto.xlsx")
    assert r.status_code == 404


def test_aggiungi_path_traversal_negato(app2, client2):
    """Finding 4: niente sorgente fuori da SOURCE_DIR in /api/aggiungi."""
    r = client2.post("/api/aggiungi", json={
        "sorgente": "../../segreto.xlsx", "etichetta": "X", "riga_etichetta": 1,
    })
    assert r.status_code == 404


def test_modella_path_traversal_negato(app2, client2):
    """Finding 4: niente scrittura fuori da DEST_DIR in /api/modella."""
    r = client2.post("/api/modella", json={
        "destinazione": "../../segreto.xlsx", "testo": "X", "riga": 1, "colonna": "A",
    })
    assert r.status_code == 404


def test_destinazione_path_traversal_negato(client2):
    """Finding 4: /api/destinazione rifiuta un nome pericoloso."""
    r = client2.post("/api/destinazione", json={"destinazione": "../../segreto.xlsx"})
    assert r.status_code == 400


def test_elabora_spazio_insufficiente_all_or_nothing(app2, client2):
    """Finding 2+7: con 5 valori ma spazio per 1, NON si compila nulla,
    lo stato NON viene salvato, e il secondo giro non dice 'nessuna variazione'."""
    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "ALFA", "riga_etichetta": 1,
    })
    client2.post("/api/aggiungi", json={
        "sorgente": "data.xlsx", "etichetta": "DATA", "riga_etichetta": 1,
    })
    client2.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})

    r1 = client2.post("/api/elabora")
    d1 = r1.get_json()
    assert r1.status_code == 200
    assert d1["ok"] is False
    assert any("spazio insufficiente" in e.get("errore", "") for e in d1["esiti"])

    # niente salvato su disco: secondo giro ripete il tentativo (non congelato)
    r2 = client2.post("/api/elabora")
    assert r2.get_json()["ok"] is False

    import openpyxl

    wb = openpyxl.load_workbook(app2.DEST_DIR / "scheda.xlsx")
    ws = wb.active
    # BETA (riga 5) NON deve essere stata svuotata
    assert ws["B5"].value == "BETA"
    wb.close()


def test_elabora_sorgente_mancante_non_svuota(client2):
    """Finding 3: sorgente sparita -> ok:false e la scheda NON viene svuotata."""
    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "ALFA", "riga_etichetta": 1,
    })
    client2.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})

    # config punta a a.xlsx; lo cancelliamo per simulare la sparizione
    import pathlib
    import app as appmod

    (appmod.SOURCE_DIR / "a.xlsx").unlink()

    # abbiamo gia' un valore in B3; deve restare
    r = client2.post("/api/elabora")
    assert r.get_json()["ok"] is False
    assert any(e.get("errore") == "sorgente mancante" for e in r.get_json()["esiti"])

    import openpyxl

    wb = openpyxl.load_workbook(appmod.DEST_DIR / "scheda.xlsx")
    assert wb.active["B3"].value == 1
    assert wb.active["B2"].value == "ALFA"
    wb.close()


def test_elabora_datetime_stato_serializzato(app2, client2):
    """Finding 5: valori data/ora NON fanno esplodere il salvataggio dello stato."""
    from datetime import datetime

    client2.post("/api/aggiungi", json={
        "sorgente": "data.xlsx", "etichetta": "DATA", "riga_etichetta": 1,
    })
    client2.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})

    r1 = client2.post("/api/elabora")
    assert r1.status_code == 200
    assert r1.get_json()["ok"] is True

    import openpyxl

    wb = openpyxl.load_workbook(app2.DEST_DIR / "scheda.xlsx")
    ws = wb.active
    assert ws["E7"].value == "DATA"
    assert ws["E8"].value == datetime(2026, 9, 1, 10, 30)

    # secondo giro: nessuna variazione (il confronto regge sui serializzati)
    r2 = client2.post("/api/elabora")
    assert r2.get_json()["ok"] is False
    assert "nessuna variazione" in r2.get_json()["motivo"]


def test_elabora_destinazione_mancante_ok_false(app2, client2):
    """Finding 6: destinazione impostata ma file sparito -> ok:false, stato non salvato."""
    import app as appmod

    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "ALFA", "riga_etichetta": 1,
    })
    client2.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})
    (appmod.DEST_DIR / "scheda.xlsx").unlink()

    r1 = client2.post("/api/elabora")
    d1 = r1.get_json()
    assert d1["ok"] is False
    assert "non trovato" in (d1.get("motivo") or "")

    # stato non salvato: ripristino il file e il giro dopo compila davvero
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["B2"] = "ALFA"
    ws["B2"].font = Font(bold=True)
    wb.save(appmod.DEST_DIR / "scheda.xlsx")

    r2 = client2.post("/api/elabora")
    assert r2.get_json()["ok"] is True


def test_export_nome_foglio_unico_per_argomento(app2, client2):
    """Finding 13 (controllo base): un argomento produce un solo foglio."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "CO2"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws["A2"] = 1
    wb.save(app2.SOURCE_DIR / "arb.xlsx")
    wb.close()

    client2.post("/api/aggiungi", json={
        "sorgente": "arb.xlsx", "etichetta": "CO2", "riga_etichetta": 1,
    })
    r = client2.post("/api/export", json={"nome": "out.xlsx"})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(app2.EXPORT_DIR / "out.xlsx")
    assert wb.sheetnames == ["CO2"]
    wb.close()


def test_export_dedup_etichette_distinte_collidono(app2, client2):
    """Finding 13: 'VEL/OCITA' e 'VEL:OCITA' collidono dopo la sanificazione
    in '_' e vengono distinte con suffisso numerico."""
    import openpyxl

    src = app2.SOURCE_DIR / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "VEL/OCITA'"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws["A2"] = 1
    ws["A4"] = "VEL:OCITA'"
    ws["A4"].font = openpyxl.styles.Font(bold=True)
    ws["A5"] = 2
    wb.save(src)
    wb.close()

    client2.post("/api/aggiungi", json={
        "sorgente": "b.xlsx", "etichetta": "VEL/OCITA'", "riga_etichetta": 1,
    })
    client2.post("/api/aggiungi", json={
        "sorgente": "b.xlsx", "etichetta": "VEL:OCITA'", "riga_etichetta": 4,
    })

    r = client2.post("/api/export", json={"nome": "out.xlsx"})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(app2.EXPORT_DIR / "out.xlsx")
    # l'apostrofo ai bordi viene tolto (regola Excel): entrambe -> VEL_OCITA, dedup
    assert wb.sheetnames[0] == "VEL_OCITA"
    assert wb.sheetnames[1] == "VEL_OCITA_2"
    wb.close()


def test_compila_senza_destinazione_400(app2, client2):
    """Finding 11 (server): compila senza scheda risponde 400 con errore."""
    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "ALFA", "riga_etichetta": 1,
    })
    r = client2.post("/api/compila")
    assert r.status_code == 400
    assert "Destinazione" in r.get_json()["errore"]


def test_elabora_incompleta_non_salva_stato_e_non_svuota(app2, client2):
    """Finding 7: seconda etichetta non in scheda -> la prima NON viene scritta."""
    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "ALFA", "riga_etichetta": 1,
    })
    # GAMMA non esiste nella scheda
    client2.post("/api/aggiungi", json={
        "sorgente": "a.xlsx", "etichetta": "GAMMA", "riga_etichetta": 1,
    })
    client2.post("/api/destinazione", json={"destinazione": "scheda.xlsx"})

    import openpyxl

    wb = openpyxl.load_workbook(app2.DEST_DIR / "scheda.xlsx")
    prima = wb.active["B3"].value
    wb.close()

    r = client2.post("/api/elabora")
    d = r.get_json()
    assert d["ok"] is False
    assert any(not e.get("trovata") for e in d["esiti"])

    # la scheda NON e' stata toccata: il vecchio valore in B3 resta
    wb = openpyxl.load_workbook(app2.DEST_DIR / "scheda.xlsx")
    assert wb.active["B3"].value == prima
    wb.close()


def test_rotta_inesistente_404_non_500(client):
    """Bug sweep B1: una rotta inesistente deve restituire 404, non 500."""
    r = client.get("/api/non-esiste")
    assert r.status_code == 404
    assert isinstance(r.get_json(), dict)


def test_metodo_non_consentito_405_non_500(client):
    """Bug sweep B1: un metodo sbagliato su una rotta esistente = 405, non 500."""
    r = client.get("/api/compila")
    assert r.status_code == 405


def test_json_malformato_400_non_500(client):
    """Bug sweep B1: body non-JSON su una rotta POST = 400, non 500."""
    r = client.post("/api/aggiungi", data="non-json", content_type="text/plain")
    assert r.status_code == 400


def test_payload_non_dict_400(client):
    """Bug sweep B2: un body JSON che non è un oggetto (lista/stringa) = 400."""
    for rotta in ("aggiungi", "rimuovi", "destinazione", "modella", "export"):
        r = client.post(f"/api/{rotta}", json=["x"])
        assert r.status_code == 400, rotta


def test_config_senza_riga_etichetta_per_voce(app2, client2):
    """Bug sweep B3: config con `riga_etichetta` mancante -> esito per-voce,
    NON un 500 (config editata a mano)."""
    import json as jsonlib

    (app2.CONFIG_PATH.parent).mkdir(parents=True, exist_ok=True)
    (app2.CONFIG_PATH).write_text(jsonlib.dumps({
        "destinazione": "scheda.xlsx",
        "argomenti": [{"sorgente": "a.xlsx", "etichetta": "ALFA"}],
    }), encoding="utf-8")

    r = client2.post("/api/compila")
    assert r.status_code == 400
    d = r.get_json()
    assert d["esiti"][0]["errore"] == "riga_etichetta non valida"


def test_export_con_sorgente_errata_400(client):
    """Bug sweep B4: export con un argomento non estraibile = 400 (tutto-o-niente),
    non un foglio con soli header."""
    client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 1,
    })
    # sorgente poi sparita
    import app as appmod

    (appmod.SOURCE_DIR / "imp.xlsx").unlink()

    r = client.post("/api/export", json={"nome": "out.xlsx"})
    assert r.status_code == 400
    assert r.get_json()["errore"] == "Nessun argomento estraibile"


def test_download_url_encoded_non_raggiunge_altri_file(client):
    """Bug sweep B6: /api/download con percorsi url-encoded non deve permettere
    di leggere file fuori da export/."""
    # nessun file esportato: qualunque tentativo deve essere 404
    for tentativo in ("..%2Fsegreto.xlsx", "..%2F..%2Fsegreto.xlsx"):
        r = client.get("/api/download/" + tentativo)
        assert r.status_code == 404, tentativo


def test_modella_su_cella_unita_400(app2, client2):
    """Bug sweep A3 (API): modella verso una cella dentro un merge = 400."""
    import openpyxl

    p = app2.DEST_DIR / "scheda.xlsx"
    wb = openpyxl.load_workbook(p)
    wb.active.merge_cells("C20:C21")
    wb.save(p)
    wb.close()

    r = client2.post("/api/modella", json={
        "destinazione": "scheda.xlsx", "testo": "X", "riga": 20, "colonna": "C",
    })
    assert r.status_code == 400
    assert "unito" in r.get_json()["errore"]


def test_modella_pulisce_spazi_testo(app2, client2):
    """Bug sweep B9: il testo di una voce viene ripulito dagli spazi di bordo."""
    r = client2.post("/api/modella", json={
        "destinazione": "scheda.xlsx", "testo": "  NOTA  ", "riga": 40, "colonna": "B",
    })
    assert r.status_code == 200
    import openpyxl

    wb = openpyxl.load_workbook(app2.DEST_DIR / "scheda.xlsx")
    assert wb.active["B40"].value == "NOTA"
    wb.close()


def test_avvio_crea_dir_config_state():
    """B2: all'avvio il modulo crea anche config/ e state/, non solo source/dest/export."""
    import app as appmod

    assert appmod.CONFIG_PATH.parent.exists()
    assert appmod.STATE_PATH.parent.exists()


def test_config_argomenti_non_lista_errore_per_voce(app2, client2):
    """B1: argomenti come dict o lista di stringhe non deve far esplodere la
    compilazione: ogni voce non valida produce un errore per-voce."""
    for config in (
        {"destinazione": "scheda.xlsx", "argomenti": {"a": "b"}},
        {"destinazione": "scheda.xlsx", "argomenti": ["mario", "luigi"]},
    ):
        app2.CONFIG_PATH.write_text(
            json.dumps(config, ensure_ascii=False), encoding="utf-8"
        )
        r = client2.post("/api/compila")
        assert r.status_code == 400
        assert any(
            e.get("errore") == "voce di configurazione non valida"
            for e in r.get_json()["esiti"]
        )


def test_sorgente_corrotto_errore_per_voce(app2, client2):
    """B3: un sorgente xlsx corrotto/illeggibile non fa cadere compila/elabora,
    ma produce un errore per-voce (e ok:false su elabora)."""
    (app2.SOURCE_DIR / "rotto.xlsx").write_bytes(b"non-e-un-zip-valido")
    app2.CONFIG_PATH.write_text(
        json.dumps({
            "destinazione": "scheda.xlsx",
            "argomenti": [{
                "sorgente": "rotto.xlsx", "etichetta": "ROTTA", "riga_etichetta": 1,
            }],
        }, ensure_ascii=False),
        encoding="utf-8",
    )
    r = client2.post("/api/compila")
    assert r.status_code == 400
    assert r.get_json()["esiti"][0]["errore"] == "sorgente non leggibile"

    re = client2.post("/api/elabora")
    assert re.status_code == 200
    assert re.get_json()["ok"] is False


def test_aggiungi_riga_frazionaria_400(client):
    """B4: una riga_etichetta frazionaria (3.5) è rifiutata; un float intero
    equivalente a un int (3.0) resta accettato."""
    r = client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 3.5,
    })
    assert r.status_code == 400

    r2 = client.post("/api/aggiungi", json={
        "sorgente": "imp.xlsx", "etichetta": "VELOCITA'", "riga_etichetta": 3.0,
    })
    assert r2.status_code == 200


def test_compila_senza_argomenti_400(client):
    """B5: compila con zero argomenti configurati risponde 400, coerente con
    elabora (che darebbe ok:false)."""
    assert client.post("/api/compila").status_code == 400
    assert client.post("/api/compila").get_json()["errore"] == "Nessun argomento configurato"


def test_export_dedup_nome_foglio_max_31(app2, client2):
    """B6: il suffisso di disambiguazione dei fogli duplicati non deve allungare
    il nome oltre 31 caratteri (openpyxl emette solo un warning, ma Excel
    può rifiutare il file)."""
    l31 = "T" * 31
    app2.CONFIG_PATH.write_text(
        json.dumps({
            "destinazione": "scheda.xlsx",
            "argomenti": [
                {"sorgente": "a.xlsx", "etichetta": l31, "riga_etichetta": 1},
                {"sorgente": "data.xlsx", "etichetta": l31 + "'", "riga_etichetta": 1},
            ],
        }, ensure_ascii=False),
        encoding="utf-8",
    )
    r = client2.post("/api/export", json={"nome": "lungo.xlsx"})
    assert r.status_code == 200

    import openpyxl

    wb = openpyxl.load_workbook(app2.EXPORT_DIR / "lungo.xlsx")
    assert max(len(ws.title) for ws in wb.worksheets) <= 31
    assert wb.worksheets[1].title == ("T" * 29) + "_2"
    wb.close()


def test_nome_file_con_doppi_punti_selezionabile(app, client):
    """B7: un file tipo ``rel..dati.xlsx`` (più punti interni) è elencato in UI
    e deve essere selezionabile, senza aprire buchi nel traversal check."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "VELOCITA'"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = 7
    wb.save(app.SOURCE_DIR / "rel..dati.xlsx")
    wb.close()

    r = client.get("/api/argomenti?file=rel..dati.xlsx")
    assert r.status_code == 200
    assert r.get_json()[0]["valori"] == [7]

    assert client.get("/api/argomenti?file=../scheda.xlsx").status_code == 404


def test_config_utf8_roundtrip(app2, client2):
    """BUG 6: la config con etichette accentate si salva e si rilegge identica
    (ensure_ascii=False + encoding utf-8, come fa l'app)."""
    config = {
        "destinazione": "scheda.xlsx",
        "argomenti": [{
            "sorgente": "a.xlsx", "etichetta": "TIPOLOGÌA ÀCCENTO", "riga_etichetta": 1,
        }],
    }
    app2.CONFIG_PATH.write_text(
        json.dumps(config, ensure_ascii=False), encoding="utf-8"
    )
    letto = json.loads(app2.CONFIG_PATH.read_text(encoding="utf-8"))
    assert letto == config
    assert letto["argomenti"][0]["etichetta"] == "TIPOLOGÌA ÀCCENTO"

    r = client2.get("/api/config")
    assert r.status_code == 200
    assert r.get_json()["argomenti"][0]["etichetta"] == "TIPOLOGÌA ÀCCENTO"