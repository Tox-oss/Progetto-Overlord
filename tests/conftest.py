import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = ROOT / "app"


@pytest.fixture()
def repo_root() -> Path:
    return ROOT


def _make_source(path: Path, blocchi: list[tuple[str, list]]) -> None:
    """Costruisce un file sorgente a 1 colonna con il formato Overlord:
    etichetta in grassetto, valori sotto, 2 righe vuote consecutive = fine blocco.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    riga = 1
    for etichetta, valori in blocchi:
        ws.cell(row=riga, column=1, value=etichetta).font = Font(bold=True)
        riga += 1
        for v in valori:
            ws.cell(row=riga, column=1, value=v)
            riga += 1
        riga += 2  # due righe vuote = fine blocco
    wb.save(path)


@pytest.fixture()
def sorgente(tmp_path: Path) -> Path:
    p = tmp_path / "sorgente.xlsx"
    _make_source(
        p,
        [
            ("VELOCITA'", [100, 102, None, 105]),
            ("TEMPERATURA", [22, 23]),
        ],
    )
    return p


def _make_scheda(path: Path, voci: dict[str, str]) -> None:
    """Costruisce una scheda con etichette sparse in grassetto."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    for coord, testo in voci.items():
        cell = ws[coord]
        cell.value = testo
        cell.font = Font(bold=True)
    wb.save(path)


@pytest.fixture()
def scheda(tmp_path: Path) -> Path:
    p = tmp_path / "scheda.xlsx"
    _make_scheda(p, {"B3": "VELOCITA'", "C12": "TEMPERATURA"})
    return p


@pytest.fixture()
def scheda_stessa_colonna(tmp_path: Path) -> Path:
    """Due etichette nella STESSA colonna senza due righe vuote in mezzo
    (scenario finding 1: lo svuotamento non deve cancellare la seconda)."""
    p = tmp_path / "scheda_samecol.xlsx"
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["B3"] = "VELOCITA'"
    ws["B3"].font = Font(bold=True)
    ws["B4"], ws["B5"] = 100, 101
    ws["B7"] = "PRESSIONE"
    ws["B7"].font = Font(bold=True)
    ws["B8"] = 9
    wb.save(p)
    wb.close()
    return p


@pytest.fixture()
def scheda_mergata(tmp_path: Path) -> Path:
    """Etichetta con una cella unita sotto (finding 9)."""
    p = tmp_path / "scheda_merge.xlsx"
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["B3"] = "VELOCITA'"
    ws["B3"].font = Font(bold=True)
    ws["B4"] = 99
    ws.merge_cells("B5:B6")
    wb.save(p)
    wb.close()
    return p


@pytest.fixture()
def sorgente_titolo_grande(tmp_path: Path) -> Path:
    """Titolo di sezione a 16pt con sotto cella NON grassetto (finding 8)."""
    p = tmp_path / "sorgente_titolo.xlsx"
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "TITOLO SEZIONE"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "non-etichetta"
    ws["A4"] = "VELOCITA'"
    ws["A4"].font = Font(bold=True)
    ws["A5"] = 1
    wb.save(p)
    wb.close()
    return p


@pytest.fixture()
def dir_files_app(tmp_path: Path):
    """Struttura cartelle files isolata (source/dest/export/config/state)."""
    files = tmp_path / "files"
    (files / "source").mkdir(parents=True)
    (files / "dest").mkdir()
    (files / "export").mkdir()
    (files / "config").mkdir()
    (files / "state").mkdir()
    return files