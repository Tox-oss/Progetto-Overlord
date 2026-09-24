import json
import os
import re
import traceback
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file

from parser import (
    _spazio_libero,
    carica_workbook,
    colonna_a_numero,
    crea_voce,
    estrai_colonna,
    trova_argomenti,
    trova_argomenti_ovunque,
    trova_cella_etichetta,
    scrivi_sotto,
)

PROJECT = Path(__file__).resolve().parent.parent
FILES = PROJECT / "files"
SOURCE_DIR = FILES / "source"
DEST_DIR = FILES / "dest"
EXPORT_DIR = FILES / "export"
CONFIG_PATH = FILES / "config" / "selezione.json"
STATE_PATH = FILES / "state" / "last_value.json"

for d in (SOURCE_DIR, DEST_DIR, EXPORT_DIR):
    d.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)


def _nome_valido(nome) -> bool:
    """True se `nome` è un file di livello singolo (niente separatori/`..`)."""
    return bool(
        nome
        and isinstance(nome, str)
        and Path(nome).name == nome
        and ".." not in nome
    )


def _risolvi_in(cartella: Path, nome) -> Path | None:
    """Risolvi `nome` dentro `cartella` (niente path traversal).

    Ritorna il Path se il file esiste dentro la cartella, altrimenti None.
    """
    if not _nome_valido(nome):
        return None
    path = (cartella / nome).resolve()
    if not path.exists() or path.parent != cartella.resolve():
        return None
    return path


@app.errorhandler(Exception)
def _errore_generico(exc):  # noqa: BLE001
    traceback.print_exc()
    return jsonify({"errore": "Errore interno del server"}), 500


def _xlsx_files(cartella: Path) -> list[str]:
    return sorted(p.name for p in cartella.glob("*.xlsx"))


def _leggi_config() -> dict:
    if CONFIG_PATH.exists():
        try:
            cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            cfg = {}
    else:
        cfg = {}
    cfg.setdefault("destinazione", None)
    cfg.setdefault("argomenti", [])
    return cfg


def _valore_seriale(v):
    """Converte un valore Excel in qualcosa che `json.dumps` sa serializzare.

    Le celle data/ora sono `datetime`/`date`, talvolta `Decimal`: senza questa
    conversione il salvataggio dello stato fa esplodere `json.dumps`.
    """
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return v


def _estratti_seriali(estratti: list[dict]) -> list[dict]:
    out = []
    for e in estratti:
        a_ = dict(e)
        a_["valori"] = [_valore_seriale(v) for v in a_.get("valori", [])]
        out.append(a_)
    return out


def _scrivi_config(payload: dict) -> None:
    tmp = CONFIG_PATH.with_suffix(".tmp")
    tmp.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    os.replace(tmp, CONFIG_PATH)


def _leggi_state() -> dict:
    if STATE_PATH.exists():
        try:
            return json.loads(STATE_PATH.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def _scrivi_state(ultimi: list) -> None:
    payload = {
        "ultimi_valori": ultimi,
        "data": datetime.now().isoformat(),
    }
    tmp = STATE_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, STATE_PATH)


def _estraai_tutti(cfg: dict) -> list[dict]:
    estratti = []
    for a in cfg.get("argomenti", []):
        if not _nome_valido(a.get("sorgente")):
            a_ = dict(a)
            a_["valori"] = []
            a_["errore"] = "sorgente non valido"
            estratti.append(a_)
            continue
        path = SOURCE_DIR / a["sorgente"]
        if not path.exists():
            a_ = dict(a)
            a_["valori"] = []
            a_["errore"] = "sorgente mancante"
            estratti.append(a_)
            continue
        a_ = dict(a)
        try:
            a_["valori"] = estrai_colonna(path, int(a_["riga_etichetta"]))
        except (TypeError, ValueError):
            a_["valori"] = []
            a_["errore"] = "riga_etichetta non valida"
        estratti.append(a_)
    return estratti


def _compila_tutti(cfg: dict, estratti: list[dict]) -> dict:
    """Compila la scheda: tutto-o-niente.

    Verifica PRIMA di scrivere che tutti gli argomenti siano estraibili
    (niente errori), trovati in scheda e che lo spazio sotto l'etichetta basti
    a contenere i valori. Al primo problema torna `{completa: False, esiti}`
    SENZA toccare la scheda. Solo se tutto è verificato scrive e salva.
    """
    if not cfg.get("destinazione"):
        return {"completa": False, "esiti": [], "errore": "Destinazione non configurata"}
    dest = _risolvi_in(DEST_DIR, cfg["destinazione"])
    if not dest:
        return {
            "completa": False,
            "esiti": [],
            "errore": "File di destinazione non trovato",
        }

    wb = carica_workbook(dest)
    ws = wb.active

    phase = []
    for e in estratti:
        voce = {
            "sorgente": e.get("sorgente"),
            "etichetta": e.get("etichetta"),
            "trovata": False,
        }
        if e.get("errore"):
            voce["errore"] = e["errore"]
            phase.append(("errore", voce))
            continue
        pos = trova_cella_etichetta(dest, e["etichetta"])
        if pos is None:
            phase.append(("non_trovata", voce))
            continue
        riga, colonna = pos
        valori = e.get("valori", [])
        if not valori:
            voce["errore"] = "nessun valore estratto dalla sorgente"
            phase.append(("spazio", voce))
            continue
        spazio = _spazio_libero(ws, riga, colonna, soglia=len(valori))
        if len(valori) > spazio:
            voce["errore"] = (
                f"spazio insufficiente ({len(valori)} valori, {spazio} righe libere)"
            )
            phase.append(("spazio", voce))
            continue
        voce.update({"trovata": True, "riga": riga, "colonna": colonna, "valori": valori})
        phase.append(("ok", voce))

    wb.close()

    completata = all(kind == "ok" for kind, _ in phase)
    if not completata:
        return {"completa": False, "esiti": [v for _, v in phase]}

    for e, (_kind, voce) in zip(estratti, phase):
        if voce.get("errore"):
            continue
        voce["scritti"] = scrivi_sotto(dest, voce["riga"], voce["colonna"], e.get("valori", []))
    return {"completa": True, "esiti": [v for _, v in phase]}


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/sorgenti")
def api_sorgenti():
    return jsonify(_xlsx_files(SOURCE_DIR))


@app.get("/api/destinazioni")
def api_destinazioni():
    return jsonify(_xlsx_files(DEST_DIR))


@app.get("/api/argomenti")
def api_argomenti():
    nome = request.args.get("file", "")
    path = _risolvi_in(SOURCE_DIR, nome)
    if not path:
        return jsonify({"errore": "File sorgente non trovato"}), 404
    argomenti = [
        {
            "etichetta": a.etichetta,
            "riga": a.riga,
            "valori": estrai_colonna(path, a.riga),
        }
        for a in trova_argomenti(path)
    ]
    return jsonify(argomenti)


@app.get("/api/scheda")
def api_scheda():
    nome = request.args.get("file", "")
    path = _risolvi_in(DEST_DIR, nome)
    if not path:
        return jsonify({"errore": "Scheda non trovata"}), 404
    voci = [
        {
            "etichetta": a.etichetta,
            "riga": a.riga,
            "colonna": a.colonna,
            "posizione": f"{_colonna_lettera(a.colonna)}{a.riga}",
        }
        for a in trova_argomenti_ovunque(path)
    ]
    return jsonify(voci)


def _colonna_lettera(n: int) -> str:
    s = ""
    while n:
        n, resto = divmod(n - 1, 26)
        s = chr(ord("A") + resto) + s
    return s


def _nome_foglio_valido(etichetta: str) -> str:
    """Nome foglio Excel ammissibile (niente caratteri vietati, max 31 char)."""
    nome = re.sub(r"[\\/*?:\[\]]", "_", str(etichetta)).strip() or "ARGOMENTO"
    return nome[:31]


def _nome_export_sicuro(nome: str) -> str:
    """Ricava un nome file valido dentro EXPORT_DIR (niente path traversal né
    caratteri illegali)."""
    base = Path(nome or "").name.strip() or "esportazione.xlsx"
    base = re.sub(r"[\\/*?:<>\"|]", "_", base)
    if not base.lower().endswith(".xlsx"):
        base += ".xlsx"
    return base


@app.get("/api/config")
def api_config():
    return jsonify(_leggi_config())


@app.post("/api/aggiungi")
def api_aggiungi():
    payload = request.get_json(force=True)
    sorgente, etichetta = payload.get("sorgente"), payload.get("etichetta")
    riga = payload.get("riga_etichetta")
    if not all((sorgente, etichetta, riga)):
        return jsonify({"errore": "Parametri mancanti"}), 400
    try:
        riga_n = int(riga)
    except (TypeError, ValueError):
        return jsonify({"errore": f"Riga non valida: {riga}"}), 400
    if riga_n < 1:
        return jsonify({"errore": "Riga non valida"}), 400
    path = _risolvi_in(SOURCE_DIR, sorgente)
    if not path:
        return jsonify({"errore": "File sorgente non trovato"}), 404
    cfg = _leggi_config()
    duplicato = any(
        a["sorgente"] == sorgente and a["etichetta"] == etichetta
        for a in cfg["argomenti"]
    )
    if not duplicato:
        cfg["argomenti"].append(
            {"sorgente": sorgente, "etichetta": etichetta, "riga_etichetta": riga_n}
        )
        _scrivi_config(cfg)
    return jsonify(cfg)


@app.post("/api/rimuovi")
def api_rimuovi():
    payload = request.get_json(force=True)
    sorgente, etichetta = payload.get("sorgente"), payload.get("etichetta")
    cfg = _leggi_config()
    cfg["argomenti"] = [
        a
        for a in cfg["argomenti"]
        if not (a["sorgente"] == sorgente and a["etichetta"] == etichetta)
    ]
    _scrivi_config(cfg)
    return jsonify(cfg)


@app.post("/api/destinazione")
def api_destinazione():
    payload = request.get_json(force=True)
    destinazione = payload.get("destinazione")
    if destinazione is not None and not _nome_valido(destinazione):
        return jsonify({"errore": "Nome scheda non valido"}), 400
    cfg = _leggi_config()
    cfg["destinazione"] = destinazione
    _scrivi_config(cfg)
    return jsonify(cfg)


@app.post("/api/modella")
def api_modella():
    payload = request.get_json(force=True)
    scheda, testo, riga = (
        payload.get("destinazione"),
        payload.get("testo"),
        payload.get("riga"),
    )
    colonna = payload.get("colonna")
    if not all((scheda, testo, riga, colonna)):
        return jsonify({"errore": "Parametri mancanti"}), 400
    path = _risolvi_in(DEST_DIR, scheda)
    if not path:
        return jsonify({"errore": "Scheda non trovata"}), 404
    try:
        riga_n = int(riga)
        colonna_n = colonna_a_numero(colonna)
    except ValueError as exc:
        return jsonify({"errore": str(exc)}), 400
    if riga_n < 1 or riga_n > 1_048_576:
        return jsonify({"errore": f"Riga fuori range: {riga_n}"}), 400
    if colonna_n < 1 or colonna_n > 16_384:
        return jsonify({"errore": f"Colonna fuori range: {colonna}"}), 400
    try:
        crea_voce(path, str(testo), riga_n, colonna_n)
    except ValueError as exc:
        return jsonify({"errore": str(exc)}), 400
    return jsonify({"ok": True})


@app.post("/api/compila")
def api_compila():
    cfg = _leggi_config()
    estratti = _estraai_tutti(cfg)
    esiti = _compila_tutti(cfg, estratti)
    if not esiti.get("completa"):
        return jsonify(
            {
                "errore": esiti.get("errore") or "Compilazione non eseguita",
                "esiti": esiti["esiti"],
            }
        ), 400
    return jsonify({"esiti": esiti["esiti"]})


@app.post("/api/export")
def api_export():
    """Salva con nome: esporta tutti gli argomenti configurati in un xlsx."""
    payload = request.get_json(force=True)
    nome = _nome_export_sicuro(payload.get("nome", ""))
    cfg = _leggi_config()
    estratti = _estraai_tutti(cfg)
    if not estratti:
        return jsonify({"errore": "Nessun argomento configurato"}), 400

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    usati: set[str] = set()
    for e in estratti:
        nome_foglio = _nome_foglio_valido(e["etichetta"])
        if nome_foglio.lower() in usati:
            suffisso = 2
            while f"{nome_foglio}_{suffisso}".lower() in usati:
                suffisso += 1
            nome_foglio = f"{nome_foglio}_{suffisso}"
        usati.add(nome_foglio.lower())
        ws = wb.create_sheet(title=nome_foglio)
        ws["A1"] = "ARGOMENTO"
        ws["B1"] = "VALORE"
        for i, v in enumerate(e.get("valori", []), start=2):
            ws.cell(row=i, column=1, value=e["etichetta"])
            ws.cell(row=i, column=2, value=v)

    dest = EXPORT_DIR / nome
    wb.save(dest)
    return jsonify({"ok": True, "file": nome, "argomenti": len(estratti)})


@app.get("/api/download/<nome_file>")
def api_download(nome_file: str):
    path = (EXPORT_DIR / nome_file).resolve()
    if not path.exists() or EXPORT_DIR not in path.parents:
        return jsonify({"errore": "File non trovato"}), 404
    return send_file(path, as_attachment=True)


@app.post("/api/elabora")
def api_elabora():
    """Bridge per l'automazione (n8n): se i valori sono cambiati, compila."""
    try:
        cfg = _leggi_config()
        if not cfg.get("argomenti") or not cfg.get("destinazione"):
            return jsonify({"ok": False, "motivo": "configurazione incompleta"}), 200
        estratti = _estraai_tutti(cfg)

        estratti_seriali = _estratti_seriali(estratti)

        stato = _leggi_state()
        if stato.get("ultimi_valori") == estratti_seriali:
            return jsonify({"ok": False, "motivo": "nessuna variazione"}), 200

        esiti = _compila_tutti(cfg, estratti)
        if not esiti.get("completa"):
            return jsonify(
                {
                    "ok": False,
                    "motivo": esiti.get("errore") or "compilazione non eseguita",
                    "esiti": esiti["esiti"],
                }
            ), 200

        _scrivi_state(estratti_seriali)
        return jsonify({"ok": True, "esiti": esiti["esiti"]})
    except Exception as exc:  # noqa: BLE001
        traceback.print_exc()
        return jsonify({"errore": str(exc)}), 500


if __name__ == "__main__":
    porta = int(os.environ.get("PORT", 8010))
    debug = os.environ.get("FLASK_DEBUG", "").lower() in {"1", "true", "yes", "on"}
    app.run(host="0.0.0.0", port=porta, debug=debug)